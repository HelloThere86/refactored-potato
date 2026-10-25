from flask import Flask, render_template, request, redirect, url_for, flash, send_file, after_this_request, Response
from models import db, BusinessUnit, Customer, Sale, Payment, Expense, ProductionBatch, HarvestLot, InventoryMovement
from datetime import datetime, timedelta
import os
import sys
import sqlite3
import tempfile
import io
import csv
import threading
import time
import signal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SUPPORTED_VEGETABLES = ["Fancy Lettuce", "Iceberg Lettuce", "Celery"]

def get_bundle_dir():
    if getattr(sys, 'frozen', False):
        if hasattr(sys, '_MEIPASS'):
            return sys._MEIPASS
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

def get_user_data_dir():
    env_override = os.environ.get('REFACTORED_POTATO_DATA_DIR')
    if env_override:
        os.makedirs(env_override, exist_ok=True)
        return env_override

    if sys.platform == 'win32':
        base = os.environ.get('LOCALAPPDATA') or os.path.expanduser('~\\AppData\\Local')
        data_dir = os.path.join(base, 'RefactoredPotato')
    else:
        data_dir = os.path.expanduser('~/.refactored_potato')

    os.makedirs(data_dir, exist_ok=True)
    return data_dir

def get_db_path():
    return os.path.join(get_user_data_dir(), 'farm.db')

bundle_dir = get_bundle_dir()
template_dir = os.path.join(bundle_dir, 'templates')
static_dir = os.path.join(bundle_dir, 'static')

app = Flask(__name__, template_folder=template_dir, static_folder=static_dir)
active_db_path = get_db_path()
clean_db_path = active_db_path.replace('\\', '/')
app.config['SQLALCHEMY_DATABASE_URI'] = f"sqlite:///{clean_db_path}"
app.config['SECRET_KEY'] = 'refactored-potato-production-secret-key-v2-0'
db.init_app(app)

def init_db(app_instance=None):
    target_app = app_instance or app
    with target_app.app_context():
        db.create_all()
        if not BusinessUnit.query.first():
            db.session.add_all([
                BusinessUnit(name="Fancy Lettuce"),
                BusinessUnit(name="Iceberg Lettuce"),
                BusinessUnit(name="Celery"),
            ])
            db.session.commit()

def get_raw_connection():
    raw_conn = db.engine.raw_connection()
    conn = raw_conn.driver_connection if hasattr(raw_conn, 'driver_connection') else raw_conn.connection
    return conn, raw_conn

def create_backup_to_file(target_path):
    source_conn, raw_conn_proxy = get_raw_connection()
    try:
        target_conn = sqlite3.connect(target_path)
        with target_conn:
            source_conn.backup(target_conn)
        target_conn.close()
    finally:
        raw_conn_proxy.close()

def restore_backup_from_file(source_path):
    db.session.remove()
    db.engine.dispose()
    target_conn, raw_conn_proxy = get_raw_connection()
    try:
        source_conn = sqlite3.connect(source_path)
        with target_conn:
            source_conn.backup(target_conn)
        source_conn.close()
    finally:
        raw_conn_proxy.close()

def validate_backup(db_path):
    conn = None
    try:
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        cursor.execute("PRAGMA integrity_check;")
        result = cursor.fetchone()
        if not result or result[0].lower() != "ok":
            return False, "Database integrity check failed. File may be corrupted."
        
        required_tables = {'business_unit', 'customer', 'sale', 'payment', 'expense', 'expense_allocations'}
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
        tables = {row[0].lower() for row in cursor.fetchall()}
        missing = required_tables - tables
        if missing:
            return False, f"Missing required tables: {', '.join(missing)}."
            
        return True, "Valid"
    except Exception as e:
        return False, f"Not a valid SQLite database file. Error: {str(e)}"
    finally:
        if conn:
            conn.close()

# ==========================================
# BACKUP & RESTORE ROUTES
# ==========================================
@app.route('/backup/download', methods=['POST'])
def download_backup():
    fd, temp_path = tempfile.mkstemp(suffix='.db')
    os.close(fd)
    try:
        create_backup_to_file(temp_path)
        @after_this_request
        def remove_file(response):
            try:
                os.remove(temp_path)
            except Exception:
                pass
            return response
        timestamp = datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
        filename = f"refactored_potato_backup_{timestamp}.db"
        return send_file(temp_path, as_attachment=True, download_name=filename)
    except Exception as e:
        flash(f"Failed to create backup: {str(e)}", "danger")
        return redirect(url_for('settings'))

@app.route('/backup/restore', methods=['POST'])
def restore_backup():
    if 'backup_file' not in request.files:
        flash('No file uploaded.', 'danger')
        return redirect(url_for('settings'))
    file = request.files['backup_file']
    if file.filename == '':
        flash('No file selected.', 'danger')
        return redirect(url_for('settings'))
        
    fd, temp_path = tempfile.mkstemp(suffix='.db')
    os.close(fd)
    file.save(temp_path)
    
    is_valid, msg = validate_backup(temp_path)
    if not is_valid:
        os.remove(temp_path)
        flash(f'Restore aborted. {msg}', 'danger')
        return redirect(url_for('settings'))
        
    safety_filename = f"farm_safety_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
    safety_path = os.path.join(get_user_data_dir(), safety_filename)
    try:
        create_backup_to_file(safety_path)
    except Exception as e:
        os.remove(temp_path)
        flash(f'Failed to create safety backup: {str(e)}', 'danger')
        return redirect(url_for('settings'))
        
    try:
        restore_backup_from_file(temp_path)
        init_db(app)
        flash('Database restored successfully!', 'success')
    except Exception as e:
        try:
            restore_backup_from_file(safety_path)
            flash(f'Restore failed, but database was safely rolled back. Error: {str(e)}', 'warning')
        except Exception:
            flash(f'CRITICAL: Restore and rollback failed! Safety backup at: {safety_path}', 'danger')
    finally:
        os.remove(temp_path)
    return redirect(url_for('settings'))

@app.route('/api/health')
def health_check():
    return {"status": "ok", "app": "Refactored Potato", "version": "2.0.0"}, 200

@app.route('/shutdown', methods=['POST'])
def shutdown():
    def terminate_process():
        time.sleep(0.8)
        os.kill(os.getpid(), signal.SIGTERM if hasattr(signal, 'SIGTERM') else signal.SIGINT)
    threading.Thread(target=terminate_process, daemon=True).start()
    return render_template('shutdown.html')

# ==========================================
# DATE FILTERING & BALANCES
# ==========================================
def parse_date_filter(request):
    filter_type = request.args.get('date_filter', 'this_month')
    start_date = None
    end_date = None
    today = datetime.utcnow().date()
    
    if filter_type == 'today':
        start_date, end_date = today, today
    elif filter_type == 'this_week':
        start_date = today - timedelta(days=today.weekday())
        end_date = today
    elif filter_type == 'this_month':
        start_date, end_date = today.replace(day=1), today
    elif filter_type == 'last_month':
        first_day_this_month = today.replace(day=1)
        end_date = first_day_this_month - timedelta(days=1)
        start_date = end_date.replace(day=1)
    elif filter_type == 'this_year':
        start_date, end_date = today.replace(month=1, day=1), today
    elif filter_type == 'custom':
        try:
            start_date = datetime.strptime(request.args.get('start_date'), '%Y-%m-%d').date()
            end_date = datetime.strptime(request.args.get('end_date'), '%Y-%m-%d').date()
        except Exception:
            pass
    return filter_type, start_date, end_date

def apply_date_filter(query, model, start_date, end_date):
    if start_date:
        query = query.filter(model.date >= datetime.combine(start_date, datetime.min.time()))
    if end_date:
        query = query.filter(model.date <= datetime.combine(end_date, datetime.max.time()))
    return query

def get_customer_balance(customer_id):
    sales = db.session.query(db.func.sum(Sale.total_amount)).filter_by(customer_id=customer_id, is_voided=False).scalar() or 0
    payments = db.session.query(db.func.sum(Payment.amount)).filter_by(customer_id=customer_id, is_voided=False).scalar() or 0
    return sales - payments

def get_total_outstanding_credit():
    total_sales = db.session.query(db.func.sum(Sale.total_amount)).filter_by(is_voided=False).scalar() or 0
    total_payments = db.session.query(db.func.sum(Payment.amount)).filter_by(is_voided=False).scalar() or 0
    return total_sales - total_payments

# ==========================================
# DASHBOARD
# ==========================================
@app.route('/')
def dashboard():
    filter_type, start_date, end_date = parse_date_filter(request)
    
    sales_q = apply_date_filter(Sale.query.filter_by(is_voided=False), Sale, start_date, end_date)
    payments_q = apply_date_filter(Payment.query.filter_by(is_voided=False), Payment, start_date, end_date)
    expenses_q = apply_date_filter(Expense.query.filter_by(is_voided=False), Expense, start_date, end_date)
    
    period_revenue = sales_q.with_entities(db.func.sum(Sale.total_amount)).scalar() or 0
    period_cash = payments_q.with_entities(db.func.sum(Payment.amount)).scalar() or 0
    period_expenses = expenses_q.with_entities(db.func.sum(Expense.amount)).scalar() or 0
    
    period_profit = period_revenue - period_expenses
    period_margin = (period_profit / period_revenue * 100) if period_revenue > 0 else 'N/A'
    outstanding_debt = get_total_outstanding_credit()

    # V2.0 Field estimate review reminders (> 14 days)
    today = datetime.utcnow().date()
    active_batches = ProductionBatch.query.filter(ProductionBatch.status.in_(['growing', 'harvesting'])).all()
    review_reminders = []
    for b in active_batches:
        days_since = (today - b.last_estimate_review_date).days
        if days_since >= 14:
            review_reminders.append({"batch": b, "days": days_since})

    # Total surplus currently on hand
    surplus_on_hand = db.session.query(db.func.sum(HarvestLot.quantity_remaining)).filter_by(is_voided=False).scalar() or 0

    sales_data = sales_q.all()
    expenses_data = expenses_q.all()
    trends = {}
    for s in sales_data:
        d = s.date.strftime('%Y-%m-%d')
        trends.setdefault(d, {'rev': 0, 'exp': 0})
        trends[d]['rev'] += s.total_amount
    for e in expenses_data:
        d = e.date.strftime('%Y-%m-%d')
        trends.setdefault(d, {'rev': 0, 'exp': 0})
        trends[d]['exp'] += e.amount
        
    chart_labels = sorted(trends.keys())
    chart_rev = [trends[d]['rev'] for d in chart_labels]
    chart_exp = [trends[d]['exp'] for d in chart_labels]
    chart_profit = [trends[d]['rev'] - trends[d]['exp'] for d in chart_labels]

    return render_template('dashboard.html', 
                           filter_type=filter_type, start_date=start_date, end_date=end_date,
                           revenue=period_revenue, cash=period_cash, debt=outstanding_debt, 
                           expenses=period_expenses, profit=period_profit, margin=period_margin,
                           chart_labels=chart_labels, chart_rev=chart_rev, chart_exp=chart_exp, chart_profit=chart_profit,
                           review_reminders=review_reminders, surplus_on_hand=surplus_on_hand)

# ==========================================
# V2.0 PRODUCTION BATCHES
# ==========================================
@app.route('/batches')
def batches():
    status_filter = request.args.get('status', 'active')
    q = ProductionBatch.query
    if status_filter == 'active':
        q = q.filter(ProductionBatch.status.in_(['growing', 'harvesting']))
    elif status_filter == 'completed':
        q = q.filter_by(status='completed')
    elif status_filter == 'all':
        pass
    all_batches = q.order_by(ProductionBatch.planting_date.desc()).all()
    
    today = datetime.utcnow().date()
    for b in all_batches:
        b.days_since_review = (today - b.last_estimate_review_date).days
        b.needs_review = b.status in ['growing', 'harvesting'] and b.days_since_review >= 14

    return render_template('batches.html', batches=all_batches, status_filter=status_filter)

@app.route('/batches/new', methods=['GET', 'POST'])
def new_batch():
    if request.method == 'POST':
        batch_code = request.form['batch_code'].strip()
        if ProductionBatch.query.filter_by(batch_code=batch_code).first():
            flash(f"Batch code '{batch_code}' is already in use. Please use a unique code.", "danger")
            units = BusinessUnit.query.filter(BusinessUnit.name.in_(SUPPORTED_VEGETABLES), BusinessUnit.is_active == True).all()
            return render_template('batch_form.html', units=units)

        unit_id = int(request.form['unit_id'])
        planting_date = datetime.strptime(request.form['planting_date'], '%Y-%m-%d').date()
        quantity_planted = int(request.form['quantity_planted'])
        planted_unit = request.form.get('planted_unit', 'seedlings').strip()
        estimated_qty = int(request.form['estimated_remaining_quantity'])
        
        expected_date = None
        if request.form.get('expected_harvest_date'):
            expected_date = datetime.strptime(request.form['expected_harvest_date'], '%Y-%m-%d').date()
            
        expected_yield = None
        if request.form.get('expected_yield'):
            expected_yield = int(request.form['expected_yield'])

        batch = ProductionBatch(
            batch_code=batch_code,
            unit_id=unit_id,
            planting_date=planting_date,
            quantity_planted=quantity_planted,
            planted_unit=planted_unit,
            estimated_remaining_quantity=estimated_qty,
            expected_harvest_date=expected_date,
            expected_yield=expected_yield,
            status='growing',
            last_estimate_review_date=planting_date,
            notes=request.form.get('notes')
        )
        db.session.add(batch)
        db.session.commit()
        flash(f"Production batch '{batch_code}' created successfully.", "success")
        return redirect(url_for('batches'))

    units = BusinessUnit.query.filter(BusinessUnit.name.in_(SUPPORTED_VEGETABLES), BusinessUnit.is_active == True).all()
    return render_template('batch_form.html', units=units)

@app.route('/batches/<int:id>')
def batch_detail(id):
    batch = ProductionBatch.query.get_or_404(id)
    harvests = HarvestLot.query.filter_by(batch_id=batch.id).order_by(HarvestLot.harvest_date.desc()).all()
    today = datetime.utcnow().date()
    days_since_review = (today - batch.last_estimate_review_date).days
    needs_review = batch.status in ['growing', 'harvesting'] and days_since_review >= 14
    return render_template('batch_detail.html', batch=batch, harvests=harvests, days_since_review=days_since_review, needs_review=needs_review)

@app.route('/batches/<int:id>/update-estimate', methods=['POST'])
def update_batch_estimate(id):
    batch = ProductionBatch.query.get_or_404(id)
    try:
        new_estimate = int(request.form['estimated_remaining_quantity'])
        if new_estimate < 0:
            flash("Estimated remaining quantity cannot be negative.", "danger")
            return redirect(url_for('batch_detail', id=batch.id))
        
        batch.estimated_remaining_quantity = new_estimate
        batch.last_estimate_review_date = datetime.utcnow().date()
        if request.form.get('status'):
            batch.status = request.form['status']
        if request.form.get('notes'):
            batch.notes = request.form['notes']

        db.session.commit()
        flash(f"Field estimate for {batch.batch_code} updated to {new_estimate} heads/stalks. Review recorded.", "success")
    except ValueError:
        flash("Invalid quantity entered.", "danger")
    return redirect(url_for('batch_detail', id=batch.id))

# ==========================================
# V2.0 HARVEST INVENTORY & WASTE LOGGING
# ==========================================
@app.route('/inventory')
def inventory():
    surplus_lots = HarvestLot.query.filter(HarvestLot.quantity_remaining > 0, HarvestLot.is_voided == False).order_by(HarvestLot.harvest_date.asc()).all()
    movements = InventoryMovement.query.order_by(InventoryMovement.timestamp.desc()).limit(30).all()
    return render_template('inventory.html', surplus_lots=surplus_lots, movements=movements)

@app.route('/inventory/waste', methods=['POST'])
def log_inventory_waste():
    lot_id = int(request.form['lot_id'])
    lot = HarvestLot.query.get_or_404(lot_id)
    waste_qty = int(request.form['quantity'])
    movement_type = request.form.get('movement_type', 'waste')
    reason = request.form.get('reason', '').strip()

    if waste_qty <= 0 or waste_qty > lot.quantity_remaining:
        flash(f"Invalid quantity. You can only adjust up to {lot.quantity_remaining} heads.", "danger")
        return redirect(url_for('inventory'))

    lot.quantity_remaining -= waste_qty
    mov = InventoryMovement(
        harvest_lot_id=lot.id,
        movement_type=movement_type,
        quantity=-waste_qty,
        reason=reason or f"Recorded {movement_type}"
    )
    db.session.add(mov)
    db.session.commit()
    flash(f"Logged {waste_qty} heads/stalks as {movement_type}. Remaining surplus updated.", "info")
    return redirect(url_for('inventory'))

# ==========================================
# COMBINED SALE & HARVEST (ATOMIC CUT-TO-ORDER)
# ==========================================
@app.route('/sales/new', methods=['GET', 'POST'])
def record_sale():
    if request.method == 'POST':
        unit = BusinessUnit.query.get_or_404(int(request.form['unit_id']))
        is_veg = unit.name in SUPPORTED_VEGETABLES

        try:
            qty_sold = float(request.form['quantity'])
            unit_price = float(request.form['unit_price'])
            if qty_sold <= 0 or unit_price < 0:
                raise ValueError("Quantity and price must be positive numbers.")
        except ValueError as e:
            flash(str(e), "danger")
            return redirect(url_for('record_sale'))

        total = qty_sold * unit_price
        batch = None
        qty_harvested = qty_sold
        customer_id = int(request.form['customer_id'])

        if is_veg:
            if not qty_sold.is_integer():
                flash("Quantity for lettuce and celery must be whole numbers (heads/stalks).", "danger")
                return redirect(url_for('record_sale'))
            qty_sold = int(qty_sold)

            batch_id = request.form.get('batch_id')
            if not batch_id:
                flash(f"Please select an active production batch for {unit.name}.", "danger")
                return redirect(url_for('record_sale'))
            
            batch = ProductionBatch.query.get_or_404(int(batch_id))
            try:
                qty_harvested = int(request.form.get('quantity_harvested', qty_sold))
            except ValueError:
                qty_harvested = qty_sold

            if qty_harvested < qty_sold:
                flash("Quantity harvested cannot be less than quantity sold to the customer.", "danger")
                return redirect(url_for('record_sale'))

            # Check field availability against harvest amount
            if qty_harvested > batch.estimated_remaining_quantity:
                flash(f"Insufficient field availability. Batch '{batch.batch_code}' only has {batch.estimated_remaining_quantity} saleable units remaining, but you entered {qty_harvested} to harvest. Please update the field estimate or select another batch.", "danger")
                return redirect(url_for('record_sale'))

        # Duplicate submission check (guard against double clicks within 3 seconds)
        recent_duplicate = Sale.query.filter(
            Sale.customer_id == customer_id,
            Sale.unit_id == unit.id,
            Sale.quantity == qty_sold,
            Sale.unit_price == unit_price,
            Sale.date >= datetime.utcnow() - timedelta(seconds=3)
        ).first()
        if recent_duplicate:
            flash("Duplicate transaction detected. Please wait a moment before resubmitting.", "warning")
            return redirect(url_for('dashboard'))

        # ATOMIC TRANSACTION EXECUTION
        try:
            new_sale = Sale(
                customer_id=customer_id,
                unit_id=unit.id,
                batch_id=batch.id if batch else None,
                quantity=qty_sold,
                unit_price=unit_price,
                total_amount=total,
                notes=request.form.get('notes')
            )
            db.session.add(new_sale)
            db.session.flush()  # obtain new_sale.id

            if is_veg and batch:
                # Deduct from field availability
                batch.estimated_remaining_quantity -= qty_harvested
                surplus = qty_harvested - qty_sold
                harvest_lot = HarvestLot(
                    batch_id=batch.id,
                    sale_id=new_sale.id,
                    quantity_harvested=qty_harvested,
                    quantity_sold=qty_sold,
                    quantity_surplus=surplus,
                    quantity_remaining=surplus,
                    notes=f"Harvested for Sale #{new_sale.id} ({qty_sold} sold, {surplus} extra)"
                )
                db.session.add(harvest_lot)
                db.session.flush()

                if surplus > 0:
                    mov = InventoryMovement(
                        harvest_lot_id=harvest_lot.id,
                        movement_type="surplus_created",
                        quantity=surplus,
                        reason=f"Extra {surplus} cut for customer order on Sale #{new_sale.id}"
                    )
                    db.session.add(mov)

            amount_paid = float(request.form.get('amount_paid', 0))
            if amount_paid > 0:
                payment = Payment(
                    customer_id=new_sale.customer_id,
                    amount=amount_paid,
                    notes="Paid at time of sale"
                )
                db.session.add(payment)

            db.session.commit()
            if is_veg and qty_harvested > qty_sold:
                flash(f"Sale recorded! Harvested {qty_harvested} units ({qty_sold} sold, {qty_harvested - qty_sold} surplus logged to cut inventory).", "success")
            else:
                flash("Sale recorded successfully.", "success")
            return redirect(url_for('dashboard'))

        except Exception as e:
            db.session.rollback()
            flash(f"Transaction failed and was completely rolled back: {str(e)}", "danger")
            return redirect(url_for('record_sale'))

    customers = Customer.query.filter_by(is_active=True).all()
    units = BusinessUnit.query.filter_by(is_active=True).all()
    active_batches = ProductionBatch.query.filter(ProductionBatch.status.in_(['growing', 'harvesting']), ProductionBatch.estimated_remaining_quantity > 0).all()
    return render_template('record_sale.html', customers=customers, units=units, batches=active_batches, veg_names=SUPPORTED_VEGETABLES)

# ==========================================
# LEDGER & SAFE SALE VOIDING
# ==========================================
@app.route('/ledger')
def ledger():
    filter_type, start_date, end_date = parse_date_filter(request)
    unit_id_filter = request.args.get('unit_id')
    customer_id_filter = request.args.get('customer_id')
    type_filter = request.args.get('type')

    sales_q = apply_date_filter(Sale.query, Sale, start_date, end_date)
    payments_q = apply_date_filter(Payment.query, Payment, start_date, end_date)
    expenses_q = apply_date_filter(Expense.query, Expense, start_date, end_date)
    
    if unit_id_filter:
        sales_q = sales_q.filter_by(unit_id=unit_id_filter)
        expenses_q = expenses_q.filter(Expense.allocated_units.any(id=unit_id_filter))
        payments_q = payments_q.filter(False)
        
    if customer_id_filter:
        sales_q = sales_q.filter_by(customer_id=customer_id_filter)
        payments_q = payments_q.filter_by(customer_id=customer_id_filter)
        expenses_q = expenses_q.filter(False)

    transactions = []
    if not type_filter or type_filter == 'Sale':
        for s in sales_q.all():
            batch_label = f" [Batch: {s.batch.batch_code}]" if s.batch else ""
            transactions.append({
                'id': s.id, 'type': 'Sale', 'date': s.date, 
                'desc': f"Customer: {s.customer.name}{batch_label}", 
                'business': s.business_unit.name, 
                'amount': s.total_amount, 'is_voided': s.is_voided
            })
            
    if not type_filter or type_filter == 'Payment':
        for p in payments_q.all():
            transactions.append({
                'id': p.id, 'type': 'Payment', 'date': p.date, 
                'desc': f"From: {p.customer.name}", 
                'business': '—', 
                'amount': p.amount, 'is_voided': p.is_voided
            })
            
    if not type_filter or type_filter == 'Expense':
        for e in expenses_q.all():
            if e.scope == 'specific': biz_label = e.allocated_units[0].name if e.allocated_units else 'Unknown'
            elif e.scope == 'shared': biz_label = f"Shared ({len(e.allocated_units)} units)"
            elif e.scope == 'farm': biz_label = "Entire Farm"
            else: biz_label = "Unallocated"
                
            transactions.append({
                'id': e.id, 'type': 'Expense', 'date': e.date, 
                'desc': e.description, 'business': biz_label, 
                'amount': e.amount, 'is_voided': e.is_voided
            })
            
    transactions.sort(key=lambda x: x['date'], reverse=True)
    units = BusinessUnit.query.filter_by(is_active=True).all()
    customers = Customer.query.filter_by(is_active=True).all()
    
    return render_template('ledger.html', 
                           transactions=transactions, units=units, customers=customers,
                           filter_type=filter_type, start_date=start_date, end_date=end_date,
                           unit_id_filter=unit_id_filter, customer_id_filter=customer_id_filter, type_filter=type_filter)

# ==========================================
# LEDGER & CORRECTED HARVEST-AWARE VOIDING
# ==========================================
@app.route('/ledger/void/<type>/<int:id>')
def void_transaction(type, id):
    if type == 'Sale':
        sale = Sale.query.get_or_404(id)
        if sale.is_voided:
            flash("Sale is already voided.", "info")
            return redirect(url_for('ledger'))
        sale.is_voided = True

        # V2.0 Corrected Rule: Physical harvests CANNOT be returned to field availability.
        # The financial sale is voided. The cut produce is retained in surplus cut inventory.
        if sale.harvest_lot and not sale.harvest_lot.is_voided:
            lot = sale.harvest_lot
            qty_cut_for_sale = int(sale.quantity)
            
            # The produce cut for this order was not sold. It becomes surplus on hand.
            lot.quantity_remaining += qty_cut_for_sale
            lot.quantity_sold = 0
            lot.notes = (lot.notes or "") + f" [Sale #{sale.id} voided - sold produce retained as surplus inventory]"

            mov = InventoryMovement(
                harvest_lot_id=lot.id,
                movement_type="sale_voided_surplus",
                quantity=qty_cut_for_sale,
                reason=f"Sale #{sale.id} voided. {qty_cut_for_sale} cut heads/stalks retained in surplus inventory. Field availability unchanged (produce was physically cut)."
            )
            db.session.add(mov)
            flash(f"Sale voided. Field availability was NOT restored because produce was physically cut; {qty_cut_for_sale} heads/stalks are now in surplus cut inventory.", "warning")
        else:
            flash("Sale voided successfully.", "warning")

    elif type == 'Payment':
        obj = Payment.query.get_or_404(id)
        obj.is_voided = True
        flash('Payment voided.', 'warning')
    elif type == 'Expense':
        obj = Expense.query.get_or_404(id)
        obj.is_voided = True
        flash('Expense voided.', 'warning')
        
    db.session.commit()
    return redirect(url_for('ledger'))

# ==========================================
# V2.0 AUDITABLE ERRONEOUS HARVEST CORRECTION
# (Crop was NOT actually cut)
# ==========================================
@app.route('/harvests/<int:id>/cancel-uncut', methods=['POST'])
def cancel_uncut_harvest(id):
    """Separate auditable correction path if a harvest entry was erroneous and crop was NEVER cut."""
    lot = HarvestLot.query.get_or_404(id)
    if lot.is_voided:
        flash("Harvest lot is already cancelled.", "info")
        return redirect(url_for('batch_detail', id=lot.batch_id))

    # Guard: Cannot cancel harvest as uncut if the financial sale is still active
    if lot.sale and not lot.sale.is_voided:
        flash(f"Cannot cancel harvest lot #{lot.id} while Sale #{lot.sale.id} is active. Please void the financial sale first.", "danger")
        return redirect(url_for('batch_detail', id=lot.batch_id))

    # Calculate what can safely be restored without creating phantom stock
    # If surplus was already wasted or consumed, we cannot restore that to the field
    available_to_restore = lot.quantity_remaining
    consumed_or_lost = lot.quantity_harvested - available_to_restore

    lot.batch.estimated_remaining_quantity += available_to_restore
    lot.quantity_remaining = 0
    lot.is_voided = True

    reason_msg = f"Harvest lot #{lot.id} cancelled as erroneous entry (crop was never physically cut). Restored {available_to_restore} to field estimate."
    if consumed_or_lost > 0:
        reason_msg += f" Note: {consumed_or_lost} heads were already consumed/wasted and could not be restored to the field."

    mov = InventoryMovement(
        harvest_lot_id=lot.id,
        movement_type="harvest_cancelled_uncut",
        quantity=-available_to_restore,
        reason=reason_msg
    )
    db.session.add(mov)
    db.session.commit()

    flash(f"Harvest lot #{lot.id} cancelled. Restored {available_to_restore} heads/stalks to {lot.batch.batch_code} field availability.", "info")
    return redirect(url_for('batch_detail', id=lot.batch_id))

# ==========================================
# PAYMENTS, EXPENSES, SETTINGS & EXPORTS
# ==========================================
@app.route('/payments/new', methods=['GET', 'POST'])
def record_payment():
    if request.method == 'POST':
        payment = Payment(
            customer_id=int(request.form['customer_id']),
            amount=float(request.form['amount']),
            notes=request.form.get('notes')
        )
        db.session.add(payment)
        db.session.commit()
        flash('Payment recorded successfully.', 'success')
        return redirect(url_for('dashboard'))
    customers = Customer.query.filter_by(is_active=True).all()
    for c in customers: c.balance = get_customer_balance(c.id)
    return render_template('record_payment.html', customers=customers)

@app.route('/expenses/new', methods=['GET', 'POST'])
def record_expense():
    if request.method == 'POST':
        scope = request.form.get('scope', 'unallocated')
        expense = Expense(
            amount=float(request.form['amount']),
            category=request.form['category'],
            description=request.form['description'],
            scope=scope
        )
        if scope == 'specific':
            unit = BusinessUnit.query.get(request.form.get('single_unit_id'))
            if unit: expense.allocated_units.append(unit)
        elif scope == 'shared':
            shared_ids = request.form.getlist('shared_unit_ids')
            for uid in shared_ids:
                unit = BusinessUnit.query.get(uid)
                if unit: expense.allocated_units.append(unit)
                
        db.session.add(expense)
        db.session.commit()
        flash('Expense recorded successfully.', 'success')
        return redirect(url_for('dashboard'))
        
    units = BusinessUnit.query.filter_by(is_active=True).all()
    categories = ["Fertilizer", "Seeds", "Feed", "Fuel", "Labour", "Transport", "Packaging", "Electricity", "Water", "Repairs", "Maintenance", "Small Purchases", "Other"]
    return render_template('record_expense.html', units=units, categories=categories)

@app.route('/settings', methods=['GET', 'POST'])
def settings():
    if request.method == 'POST':
        action = request.form.get('action')
        if action == 'add_customer':
            new_cust = Customer(name=request.form['name'], phone=request.form.get('phone'), location=request.form.get('location'), notes=request.form.get('notes'))
            db.session.add(new_cust)
            flash('Customer added.', 'success')
        elif action == 'add_unit':
            new_unit = BusinessUnit(name=request.form['name'])
            db.session.add(new_unit)
            flash('Business Unit added.', 'success')
        db.session.commit()
        return redirect(url_for('settings'))
    customers = Customer.query.all()
    for c in customers: c.balance = get_customer_balance(c.id)
    units = BusinessUnit.query.all()
    return render_template('settings.html', customers=customers, units=units, db_path=get_db_path(), data_dir=get_user_data_dir())

@app.route('/settings/edit_customer/<int:id>', methods=['GET', 'POST'])
def edit_customer(id):
    customer = Customer.query.get_or_404(id)
    if request.method == 'POST':
        customer.name = request.form['name']
        customer.phone = request.form.get('phone')
        customer.location = request.form.get('location')
        customer.notes = request.form.get('notes')
        db.session.commit()
        flash('Customer updated.', 'success')
        return redirect(url_for('settings'))
    return render_template('edit_customer.html', customer=customer)

@app.route('/settings/toggle/<type>/<int:id>')
def toggle_active(type, id):
    obj = Customer.query.get_or_404(id) if type == 'customer' else BusinessUnit.query.get_or_404(id)
    obj.is_active = not obj.is_active
    db.session.commit()
    flash('Status updated.', 'info')
    return redirect(url_for('settings'))

@app.route('/analytics')
@app.route('/profitability')
def analytics():
    filter_type, start_date, end_date = parse_date_filter(request)
    sales_q = apply_date_filter(Sale.query.filter_by(is_voided=False), Sale, start_date, end_date)
    expenses_q = apply_date_filter(Expense.query.filter_by(is_voided=False), Expense, start_date, end_date)
    
    farm_revenue = sales_q.with_entities(db.func.sum(Sale.total_amount)).scalar() or 0
    farm_expenses = expenses_q.with_entities(db.func.sum(Expense.amount)).scalar() or 0
    farm_profit = farm_revenue - farm_expenses
    farm_margin = (farm_profit / farm_revenue * 100) if farm_revenue > 0 else 'N/A'
    
    farm_direct = expenses_q.filter(Expense.scope == 'specific').with_entities(db.func.sum(Expense.amount)).scalar() or 0
    farm_shared = expenses_q.filter(Expense.scope == 'shared').with_entities(db.func.sum(Expense.amount)).scalar() or 0
    farm_wide = expenses_q.filter(Expense.scope == 'farm').with_entities(db.func.sum(Expense.amount)).scalar() or 0
    farm_unallocated = expenses_q.filter(Expense.scope == 'unallocated').with_entities(db.func.sum(Expense.amount)).scalar() or 0
    
    expenses_data = expenses_q.all()
    categories_dict = {}
    for e in expenses_data:
        categories_dict[e.category] = categories_dict.get(e.category, 0) + e.amount
    expenses_by_category = sorted(categories_dict.items(), key=lambda x: x[1], reverse=True)

    units = BusinessUnit.query.filter_by(is_active=True).all()
    performance = []
    for unit in units:
        unit_sales_q = sales_q.filter_by(unit_id=unit.id)
        direct_costs = expenses_q.filter(Expense.allocated_units.any(id=unit.id), Expense.scope == 'specific').with_entities(db.func.sum(Expense.amount)).scalar() or 0
        shared_costs_memo = expenses_q.filter(Expense.allocated_units.any(id=unit.id), Expense.scope == 'shared').with_entities(db.func.sum(Expense.amount)).scalar() or 0
        rev = unit_sales_q.with_entities(db.func.sum(Sale.total_amount)).scalar() or 0
        count = unit_sales_q.count()
        profit = rev - direct_costs
        margin = (profit / rev * 100) if rev > 0 else 'N/A'
        performance.append({'name': unit.name, 'revenue': rev, 'direct_costs': direct_costs, 'shared_costs_memo': shared_costs_memo, 'profit': profit, 'margin': margin, 'count': count})
        
    customers = Customer.query.filter_by(is_active=True).all()
    customer_analytics = []
    for c in customers:
        balance = get_customer_balance(c.id)
        if balance > 0:
            recent_sale = Sale.query.filter_by(customer_id=c.id, is_voided=False).order_by(Sale.date.desc()).first()
            recent_payment = Payment.query.filter_by(customer_id=c.id, is_voided=False).order_by(Payment.date.desc()).first()
            customer_analytics.append({'name': c.name, 'balance': balance, 'recent_sale': recent_sale.date if recent_sale else None, 'recent_payment': recent_payment.date if recent_payment else None})
    customer_analytics.sort(key=lambda x: x['balance'], reverse=True)
        
    return render_template('analytics.html', 
                           filter_type=filter_type, start_date=start_date, end_date=end_date,
                           performance=performance, farm_revenue=farm_revenue, farm_expenses=farm_expenses,
                           farm_profit=farm_profit, farm_margin=farm_margin, farm_direct=farm_direct, 
                           farm_shared=farm_shared, farm_wide=farm_wide, farm_unallocated=farm_unallocated,
                           expenses_by_category=expenses_by_category, customer_analytics=customer_analytics)

# ==========================================
# EXPORTS (CSV & EXCEL)
# ==========================================
def sanitize_for_spreadsheet(value):
    if value is None: return ""
    str_val = str(value)
    if str_val and str_val[0] in ('=', '+', '-', '@', '\t', '\r'):
        return f"'{str_val}"
    return str_val

def get_export_queries(start_date, end_date, unit_id_filter=None):
    sales_q = apply_date_filter(Sale.query, Sale, start_date, end_date)
    expenses_q = apply_date_filter(Expense.query, Expense, start_date, end_date)
    payments_q = apply_date_filter(Payment.query, Payment, start_date, end_date)
    if unit_id_filter:
        sales_q = sales_q.filter_by(unit_id=unit_id_filter)
        expenses_q = expenses_q.filter(Expense.allocated_units.any(id=unit_id_filter))
    return sales_q, expenses_q, payments_q

@app.route('/export')
def export_hub():
    filter_type, start_date, end_date = parse_date_filter(request)
    unit_id_filter = request.args.get('unit_id')
    units = BusinessUnit.query.filter_by(is_active=True).all()
    return render_template('export.html', filter_type=filter_type, start_date=start_date, end_date=end_date, unit_id_filter=unit_id_filter, units=units)

@app.route('/export/csv/<dataset>')
def export_csv(dataset):
    filter_type, start_date, end_date = parse_date_filter(request)
    unit_id_filter = request.args.get('unit_id')
    sales_q, expenses_q, payments_q = get_export_queries(start_date, end_date, unit_id_filter)
    output = io.StringIO()
    writer = csv.writer(output)
    timestamp = datetime.utcnow().strftime('%Y-%m-%d_%H%M')

    if dataset == 'sales':
        writer.writerow(['Sale ID', 'Date', 'Customer', 'Product / Business Unit', 'Quantity', 'Unit Price ($)', 'Total Amount ($)', 'Batch', 'Status', 'Notes'])
        for s in sales_q.order_by(Sale.date.desc()).all():
            writer.writerow([
                s.id, s.date.strftime('%Y-%m-%d %H:%M'),
                sanitize_for_spreadsheet(s.customer.name if s.customer else 'Unknown'),
                sanitize_for_spreadsheet(s.business_unit.name if s.business_unit else 'Unknown'),
                f"{s.quantity:.2f}", f"{s.unit_price:.2f}", f"{s.total_amount:.2f}",
                sanitize_for_spreadsheet(s.batch.batch_code if s.batch else 'N/A'),
                'VOIDED' if s.is_voided else 'Active', sanitize_for_spreadsheet(s.notes or '')
            ])
        filename = f"sales_export_{timestamp}.csv"
    elif dataset == 'expenses':
        writer.writerow(['Expense ID', 'Date', 'Category', 'Description', 'Scope', 'Allocated Business Units', 'Amount ($)', 'Status'])
        for e in expenses_q.order_by(Expense.date.desc()).all():
            units_str = e.allocated_units[0].name if e.scope == 'specific' and e.allocated_units else (", ".join([u.name for u in e.allocated_units]) if e.scope == 'shared' else ("Entire Farm" if e.scope == 'farm' else "Unallocated"))
            writer.writerow([e.id, e.date.strftime('%Y-%m-%d %H:%M'), sanitize_for_spreadsheet(e.category), sanitize_for_spreadsheet(e.description or ''), e.scope.capitalize(), sanitize_for_spreadsheet(units_str), f"{e.amount:.2f}", 'VOIDED' if e.is_voided else 'Active'])
        filename = f"expenses_export_{timestamp}.csv"
    elif dataset == 'payments':
        writer.writerow(['Payment ID', 'Date', 'Customer', 'Amount ($)', 'Status', 'Notes'])
        for p in payments_q.order_by(Payment.date.desc()).all():
            writer.writerow([p.id, p.date.strftime('%Y-%m-%d %H:%M'), sanitize_for_spreadsheet(p.customer.name if p.customer else 'Unknown'), f"{p.amount:.2f}", 'VOIDED' if p.is_voided else 'Active', sanitize_for_spreadsheet(p.notes or '')])
        filename = f"payments_export_{timestamp}.csv"
    elif dataset == 'customer_balances':
        writer.writerow(['Customer ID', 'Customer Name', 'Phone', 'Location', 'Status', 'All-Time Sales ($)', 'All-Time Payments ($)', 'Current Outstanding Balance ($)'])
        for c in Customer.query.order_by(Customer.name).all():
            all_sales = db.session.query(db.func.sum(Sale.total_amount)).filter_by(customer_id=c.id, is_voided=False).scalar() or 0.0
            all_pmts = db.session.query(db.func.sum(Payment.amount)).filter_by(customer_id=c.id, is_voided=False).scalar() or 0.0
            writer.writerow([c.id, sanitize_for_spreadsheet(c.name), sanitize_for_spreadsheet(c.phone or ''), sanitize_for_spreadsheet(c.location or ''), 'Active' if c.is_active else 'Archived', f"{all_sales:.2f}", f"{all_pmts:.2f}", f"{(all_sales - all_pmts):.2f}"])
        filename = f"customer_balances_all_time_{timestamp}.csv"
    elif dataset == 'business_summaries':
        writer.writerow(['Business Unit', 'Sales Count', 'Revenue ($)', 'Direct Expenses ($)', 'Estimated Direct Profit ($)', 'Direct Margin (%)', 'Shared Expenses Memo ($)'])
        for unit in BusinessUnit.query.filter_by(is_active=True).all():
            u_sales_q = apply_date_filter(Sale.query.filter_by(unit_id=unit.id, is_voided=False), Sale, start_date, end_date)
            u_direct_exp_q = apply_date_filter(Expense.query.filter(Expense.allocated_units.any(id=unit.id), Expense.scope == 'specific', Expense.is_voided == False), Expense, start_date, end_date)
            u_shared_exp_q = apply_date_filter(Expense.query.filter(Expense.allocated_units.any(id=unit.id), Expense.scope == 'shared', Expense.is_voided == False), Expense, start_date, end_date)
            rev = u_sales_q.with_entities(db.func.sum(Sale.total_amount)).scalar() or 0.0
            direct_costs = u_direct_exp_q.with_entities(db.func.sum(Expense.amount)).scalar() or 0.0
            shared_memo = u_shared_exp_q.with_entities(db.func.sum(Expense.amount)).scalar() or 0.0
            profit = rev - direct_costs
            margin_str = f"{(profit / rev * 100):.1f}" if rev > 0 else 'N/A'
            writer.writerow([sanitize_for_spreadsheet(unit.name), u_sales_q.count(), f"{rev:.2f}", f"{direct_costs:.2f}", f"{profit:.2f}", margin_str, f"{shared_memo:.2f}"])
        filename = f"business_unit_summary_{timestamp}.csv"
    else:
        flash("Invalid export dataset selected.", "danger")
        return redirect(url_for('export_hub'))

    bytes_data = output.getvalue().encode('utf-8-sig')
    return Response(bytes_data, mimetype="text/csv", headers={"Content-Disposition": f"attachment; filename={filename}"})

@app.route('/export/excel')
def export_excel():
    filter_type, start_date, end_date = parse_date_filter(request)
    unit_id_filter = request.args.get('unit_id')
    sales_q, expenses_q, payments_q = get_export_queries(start_date, end_date, unit_id_filter)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2C523D", end_color="2C523D", fill_type="solid")
    meta_font = Font(name="Calibri", size=10, italic=True, color="555555")
    thin_border = Border(left=Side(style='thin', color='E0E0E0'), right=Side(style='thin', color='E0E0E0'), top=Side(style='thin', color='E0E0E0'), bottom=Side(style='thin', color='E0E0E0'))

    def style_sheet(ws, title, meta_text, headers, rows):
        ws.title = title
        ws.append([f"Refactored Potato — {title}"])
        ws.cell(1, 1).font = Font(name="Calibri", size=14, bold=True, color="1E392A")
        ws.append([meta_text])
        ws.cell(2, 1).font = meta_font
        ws.append([])
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for r_data in rows:
            ws.append(r_data)
            current_row = ws.max_row
            for c_idx in range(1, len(r_data) + 1):
                c = ws.cell(row=current_row, column=c_idx)
                c.border = thin_border
                if isinstance(r_data[c_idx - 1], (int, float)):
                    c.number_format = '#,##0.00'
                    c.alignment = Alignment(horizontal="right")
        ws.freeze_panes = "A5"
        for col in ws.columns:
            max_len = max([len(str(cell.value or '')) for cell in col if cell.row >= 3] + [10])
            ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 3

    date_label = f"Period: {filter_type.replace('_', ' ').capitalize()} ({start_date or 'Beginning'} to {end_date or 'Present'})"
    sales_rows = [[s.id, s.date.strftime('%Y-%m-%d %H:%M'), sanitize_for_spreadsheet(s.customer.name if s.customer else 'Unknown'), sanitize_for_spreadsheet(s.business_unit.name if s.business_unit else 'Unknown'), s.quantity, s.unit_price, s.total_amount, 'VOIDED' if s.is_voided else 'Active', sanitize_for_spreadsheet(s.notes or '')] for s in sales_q.order_by(Sale.date.desc()).all()]
    style_sheet(wb.create_sheet(), "Sales", f"{date_label} | Voided sales excluded from totals.", ['Sale ID', 'Date', 'Customer', 'Product', 'Quantity', 'Unit Price ($)', 'Total ($)', 'Status', 'Notes'], sales_rows)

    exp_rows = [[e.id, e.date.strftime('%Y-%m-%d %H:%M'), sanitize_for_spreadsheet(e.category), sanitize_for_spreadsheet(e.description or ''), e.scope.capitalize(), sanitize_for_spreadsheet(e.allocated_units[0].name if e.scope == 'specific' and e.allocated_units else (", ".join([u.name for u in e.allocated_units]) if e.scope == 'shared' else ("Entire Farm" if e.scope == 'farm' else "Unallocated"))), e.amount, 'VOIDED' if e.is_voided else 'Active'] for e in expenses_q.order_by(Expense.date.desc()).all()]
    style_sheet(wb.create_sheet(), "Expenses", f"{date_label} | Voided expenses excluded from summaries.", ['Expense ID', 'Date', 'Category', 'Description', 'Scope', 'Allocated Units', 'Amount ($)', 'Status'], exp_rows)

    pmt_rows = [[p.id, p.date.strftime('%Y-%m-%d %H:%M'), sanitize_for_spreadsheet(p.customer.name if p.customer else 'Unknown'), p.amount, 'VOIDED' if p.is_voided else 'Active', sanitize_for_spreadsheet(p.notes or '')] for p in payments_q.order_by(Payment.date.desc()).all()]
    style_sheet(wb.create_sheet(), "Payments", f"{date_label} | Cash collections.", ['Payment ID', 'Date', 'Customer', 'Amount ($)', 'Status', 'Notes'], pmt_rows)

    cust_rows = []
    for c in Customer.query.order_by(Customer.name).all():
        s_tot = db.session.query(db.func.sum(Sale.total_amount)).filter_by(customer_id=c.id, is_voided=False).scalar() or 0.0
        p_tot = db.session.query(db.func.sum(Payment.amount)).filter_by(customer_id=c.id, is_voided=False).scalar() or 0.0
        cust_rows.append([c.id, sanitize_for_spreadsheet(c.name), sanitize_for_spreadsheet(c.phone or ''), sanitize_for_spreadsheet(c.location or ''), 'Active' if c.is_active else 'Archived', s_tot, p_tot, s_tot - p_tot])
    style_sheet(wb.create_sheet(), "Customer Balances", "All-Time Accounts Ledger | Balances reflect complete customer history.", ['Customer ID', 'Customer Name', 'Phone', 'Location', 'Status', 'All-Time Sales ($)', 'All-Time Payments ($)', 'Current Balance ($)'], cust_rows)

    bu_rows = []
    for unit in BusinessUnit.query.filter_by(is_active=True).all():
        u_sales_q = apply_date_filter(Sale.query.filter_by(unit_id=unit.id, is_voided=False), Sale, start_date, end_date)
        u_direct_exp_q = apply_date_filter(Expense.query.filter(Expense.allocated_units.any(id=unit.id), Expense.scope == 'specific', Expense.is_voided == False), Expense, start_date, end_date)
        u_shared_exp_q = apply_date_filter(Expense.query.filter(Expense.allocated_units.any(id=unit.id), Expense.scope == 'shared', Expense.is_voided == False), Expense, start_date, end_date)
        rev = u_sales_q.with_entities(db.func.sum(Sale.total_amount)).scalar() or 0.0
        direct_costs = u_direct_exp_q.with_entities(db.func.sum(Expense.amount)).scalar() or 0.0
        shared_memo = u_shared_exp_q.with_entities(db.func.sum(Expense.amount)).scalar() or 0.0
        profit = rev - direct_costs
        margin_str = f"{(profit / rev * 100):.1f}%" if rev > 0 else 'N/A'
        bu_rows.append([sanitize_for_spreadsheet(unit.name), u_sales_q.count(), rev, direct_costs, profit, margin_str, shared_memo])
    style_sheet(wb.create_sheet(), "Business Unit Summary", f"{date_label} | Direct profit excludes shared costs to prevent double counting.", ['Business Unit', 'Sales Count', 'Revenue ($)', 'Direct Costs ($)', 'Direct Profit ($)', 'Margin (%)', 'Shared Memo ($)'], bu_rows)

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    timestamp = datetime.utcnow().strftime('%Y-%m-%d_%H%M')
    return send_file(excel_buffer, as_attachment=True, download_name=f"refactored_potato_financials_{timestamp}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == '__main__':
    init_db(app)
    app.run(debug=True, port=5000)