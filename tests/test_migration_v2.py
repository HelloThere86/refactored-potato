import pytest
import os
import sqlite3
import tempfile
import io
import openpyxl
from migrate_v2_0 import migrate
from models import db, BusinessUnit, Customer, Sale, Payment, Expense
from app import get_customer_balance, validate_backup

@pytest.fixture
def v133_db_path(tmp_path):
    """Constructs a genuine physical V1.3.3 SQLite database on disk with financial records."""
    db_file = str(tmp_path / "v133_test_farm.db")
    conn = sqlite3.connect(db_file)
    c = conn.cursor()

    # V1.3.3 Schema (Sale has NO batch_id, no production tables)
    c.executescript("""
    CREATE TABLE business_unit (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name VARCHAR(100) NOT NULL,
        is_active BOOLEAN DEFAULT 1
    );

    CREATE TABLE customer (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name VARCHAR(100) NOT NULL,
        phone VARCHAR(50),
        location VARCHAR(255),
        notes TEXT,
        is_active BOOLEAN DEFAULT 1
    );

    CREATE TABLE sale (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date DATETIME DEFAULT CURRENT_TIMESTAMP,
        customer_id INTEGER NOT NULL REFERENCES customer(id),
        unit_id INTEGER NOT NULL REFERENCES business_unit(id),
        quantity FLOAT NOT NULL,
        unit_price FLOAT NOT NULL,
        total_amount FLOAT NOT NULL,
        notes VARCHAR(255),
        is_voided BOOLEAN DEFAULT 0
    );

    CREATE TABLE payment (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date DATETIME DEFAULT CURRENT_TIMESTAMP,
        customer_id INTEGER NOT NULL REFERENCES customer(id),
        amount FLOAT NOT NULL,
        notes VARCHAR(255),
        is_voided BOOLEAN DEFAULT 0
    );

    CREATE TABLE expense (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date DATETIME DEFAULT CURRENT_TIMESTAMP,
        amount FLOAT NOT NULL,
        category VARCHAR(50) NOT NULL,
        description VARCHAR(255),
        unit_id INTEGER REFERENCES business_unit(id),
        scope VARCHAR(20) DEFAULT 'unallocated',
        is_voided BOOLEAN DEFAULT 0
    );

    CREATE TABLE expense_allocations (
        expense_id INTEGER REFERENCES expense(id),
        business_unit_id INTEGER REFERENCES business_unit(id),
        PRIMARY KEY (expense_id, business_unit_id)
    );
    """)

    # Populate V1.3.3 Data
    c.execute("INSERT INTO business_unit (id, name, is_active) VALUES (1, 'Fancy Lettuce', 1);")
    c.execute("INSERT INTO business_unit (id, name, is_active) VALUES (2, 'Celery', 1);")
    c.execute("INSERT INTO business_unit (id, name, is_active) VALUES (3, 'Honey', 1);")

    c.execute("INSERT INTO customer (id, name, phone, location) VALUES (1, 'Restaurant ABC', '077123456', 'Harare');")
    c.execute("INSERT INTO customer (id, name, phone, location) VALUES (2, 'Supermarket XYZ', '077654321', 'Bulawayo');")

    # Sales (1 active $50, 1 active $50, 1 voided $10)
    c.execute("INSERT INTO sale (id, customer_id, unit_id, quantity, unit_price, total_amount, is_voided) VALUES (1, 1, 1, 100, 0.50, 50.0, 0);")
    c.execute("INSERT INTO sale (id, customer_id, unit_id, quantity, unit_price, total_amount, is_voided) VALUES (2, 2, 2, 40, 1.25, 50.0, 0);")
    c.execute("INSERT INTO sale (id, customer_id, unit_id, quantity, unit_price, total_amount, is_voided) VALUES (3, 1, 1, 20, 0.50, 10.0, 1);")

    # Payments ($30 for Cust 1, $50 for Cust 2)
    c.execute("INSERT INTO payment (id, customer_id, amount, is_voided) VALUES (1, 1, 30.0, 0);")
    c.execute("INSERT INTO payment (id, customer_id, amount, is_voided) VALUES (2, 2, 50.0, 0);")

    # Expenses across all 4 scopes
    c.execute("INSERT INTO expense (id, amount, category, description, scope, is_voided) VALUES (1, 20.0, 'Fertilizer', 'Lettuce fertilizer', 'specific', 0);")
    c.execute("INSERT INTO expense_allocations (expense_id, business_unit_id) VALUES (1, 1);")

    c.execute("INSERT INTO expense (id, amount, category, description, scope, is_voided) VALUES (2, 30.0, 'Transport', 'Shared fuel', 'shared', 0);")
    c.execute("INSERT INTO expense_allocations (expense_id, business_unit_id) VALUES (2, 1);")
    c.execute("INSERT INTO expense_allocations (expense_id, business_unit_id) VALUES (2, 2);")

    c.execute("INSERT INTO expense (id, amount, category, description, scope, is_voided) VALUES (3, 15.0, 'Electricity', 'Farm power', 'farm', 0);")
    c.execute("INSERT INTO expense (id, amount, category, description, scope, is_voided) VALUES (4, 5.0, 'Other', 'Misc small item', 'unallocated', 0);")

    conn.commit()
    conn.close()

    return db_file

def test_v133_to_v20_migration_safety(v133_db_path):
    """Verifies running V2.0 migration on a physical V1.3.3 database preserves financial records and calculations."""
    # 1. Pre-migration check: verify sale table lacks batch_id
    conn = sqlite3.connect(v133_db_path)
    c = conn.cursor()
    c.execute("PRAGMA table_info(sale);")
    cols_pre = [row[1] for row in c.fetchall()]
    assert 'batch_id' not in cols_pre
    conn.close()

    # 2. Run the migration script
    migrate(v133_db_path)

    # 3. Post-migration check: verify schema additions
    conn = sqlite3.connect(v133_db_path)
    c = conn.cursor()
    c.execute("PRAGMA table_info(sale);")
    cols_post = [row[1] for row in c.fetchall()]
    assert 'batch_id' in cols_post

    c.execute("SELECT name FROM sqlite_master WHERE type='table';")
    tables = {r[0] for r in c.fetchall()}
    assert {'production_batch', 'harvest_lot', 'inventory_movement'}.issubset(tables)

    # 4. Verify record counts are 100% preserved
    c.execute("SELECT COUNT(*) FROM sale;")
    assert c.fetchone()[0] == 3
    c.execute("SELECT COUNT(*) FROM payment;")
    assert c.fetchone()[0] == 2
    c.execute("SELECT COUNT(*) FROM expense;")
    assert c.fetchone()[0] == 4
    c.execute("SELECT COUNT(*) FROM customer;")
    assert c.fetchone()[0] == 2
    c.execute("SELECT COUNT(*) FROM expense_allocations;")
    assert c.fetchone()[0] == 3

    # 5. Verify financial totals & balances directly from the migrated file
    # Customer 1: $50 active sale, $10 voided sale (ignored), $30 payment -> Balance = $20.00
    c.execute("SELECT SUM(total_amount) FROM sale WHERE customer_id=1 AND is_voided=0;")
    sales_c1 = c.fetchone()[0] or 0.0
    c.execute("SELECT SUM(amount) FROM payment WHERE customer_id=1 AND is_voided=0;")
    pmts_c1 = c.fetchone()[0] or 0.0
    assert sales_c1 - pmts_c1 == 20.0

    # Customer 2: $50 active sale, $50 payment -> Balance = $0.00
    c.execute("SELECT SUM(total_amount) FROM sale WHERE customer_id=2 AND is_voided=0;")
    sales_c2 = c.fetchone()[0] or 0.0
    c.execute("SELECT SUM(amount) FROM payment WHERE customer_id=2 AND is_voided=0;")
    pmts_c2 = c.fetchone()[0] or 0.0
    assert sales_c2 - pmts_c2 == 0.0

    # Verify all historical sales have batch_id IS NULL
    c.execute("SELECT COUNT(*) FROM sale WHERE batch_id IS NOT NULL;")
    assert c.fetchone()[0] == 0

    # Verify all 4 expense scopes survive migration
    c.execute("SELECT DISTINCT scope FROM expense;")
    scopes = {r[0] for r in c.fetchall()}
    assert scopes == {'specific', 'shared', 'farm', 'unallocated'}

    conn.close()

    # 6. Verify V1.3.1 backup validator recognizes the migrated database
    is_valid, msg = validate_backup(v133_db_path)
    assert is_valid is True

def test_v20_exports_and_queries_with_historical_unbatched_data(client, app):
    """Verifies CSV and Excel exports handle historical sales (batch_id is None) and all 4 expense scopes."""
    with app.app_context():
        # Seed V1.3.3 historical data directly in the test database
        lettuce = BusinessUnit(name="Fancy Lettuce")
        celery = BusinessUnit(name="Celery")
        honey = BusinessUnit(name="Honey")
        cust1 = Customer(name="Restaurant ABC", phone="077123456", location="Harare")
        cust2 = Customer(name="Supermarket XYZ", phone="077654321", location="Bulawayo")
        db.session.add_all([lettuce, celery, honey, cust1, cust2])
        db.session.commit()

        # Historical sales (batch_id is None)
        sale1 = Sale(customer_id=cust1.id, unit_id=lettuce.id, quantity=100.0, unit_price=0.50, total_amount=50.0, batch_id=None, is_voided=False)
        sale2 = Sale(customer_id=cust2.id, unit_id=celery.id, quantity=40.0, unit_price=1.25, total_amount=50.0, batch_id=None, is_voided=False)
        sale3 = Sale(customer_id=cust1.id, unit_id=lettuce.id, quantity=20.0, unit_price=0.50, total_amount=10.0, batch_id=None, is_voided=True)
        db.session.add_all([sale1, sale2, sale3])

        # Payments
        pmt1 = Payment(customer_id=cust1.id, amount=30.0, is_voided=False)
        pmt2 = Payment(customer_id=cust2.id, amount=50.0, is_voided=False)
        db.session.add_all([pmt1, pmt2])

        # Expenses across all 4 scopes
        exp_specific = Expense(amount=20.0, category="Fertilizer", description="Lettuce fertilizer", scope="specific", is_voided=False)
        exp_specific.allocated_units.append(lettuce)

        exp_shared = Expense(amount=30.0, category="Transport", description="Shared fuel", scope="shared", is_voided=False)
        exp_shared.allocated_units.extend([lettuce, celery])

        exp_farm = Expense(amount=15.0, category="Electricity", description="Farm power", scope="farm", is_voided=False)
        exp_unallocated = Expense(amount=5.0, category="Other", description="Misc small item", scope="unallocated", is_voided=False)
        db.session.add_all([exp_specific, exp_shared, exp_farm, exp_unallocated])
        db.session.commit()

        # Verify customer balances calculate correctly
        assert get_customer_balance(cust1.id) == 20.0
        assert get_customer_balance(cust2.id) == 0.0

    # Test Sales CSV export with historical data (outputs 'N/A' for batch)
    res_sales = client.get("/export/csv/sales?date_filter=all_time")
    assert res_sales.status_code == 200
    assert b"Restaurant ABC" in res_sales.data
    assert b"N/A" in res_sales.data

    # Test Excel export
    res_excel = client.get("/export/excel?date_filter=all_time")
    assert res_excel.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res_excel.data))
    assert "Sales" in wb.sheetnames
    assert "Customer Balances" in wb.sheetnames
    assert "Business Unit Summary" in wb.sheetnames