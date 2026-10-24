import pytest
import io
import csv
import openpyxl
from datetime import datetime, timedelta
from models import db, BusinessUnit, Customer, Sale, Payment, Expense

@pytest.fixture
def export_test_data(app):
    """Seeds predictable multi-unit data with sales, voided items, and formula injection strings."""
    with app.app_context():
        lettuce = BusinessUnit(name="Fancy Lettuce")
        celery = BusinessUnit(name="Celery")
        db.session.add_all([lettuce, celery])
        db.session.commit()

        # Capture IDs as integers while the database session is open
        lettuce_id = lettuce.id
        celery_id = celery.id

        # Customer with formula injection test string
        cust_normal = Customer(name="Local Grocer", phone="077111222")
        cust_malicious = Customer(name="=SUM(A1:A10)", phone="+26377000000", notes="@malicious_link")
        db.session.add_all([cust_normal, cust_malicious])
        db.session.commit()

        # Sales: 1 Active, 1 Voided, 1 Other Unit
        sale_active = Sale(customer_id=cust_normal.id, unit_id=lettuce_id, quantity=10, unit_price=2.0, total_amount=20.0, is_voided=False)
        sale_voided = Sale(customer_id=cust_normal.id, unit_id=lettuce_id, quantity=5, unit_price=2.0, total_amount=10.0, is_voided=True)
        sale_celery = Sale(customer_id=cust_malicious.id, unit_id=celery_id, quantity=4, unit_price=5.0, total_amount=20.0, is_voided=False)

        # Payments: 1 Active, 1 Voided
        pmt_active = Payment(customer_id=cust_normal.id, amount=15.0, is_voided=False)
        pmt_voided = Payment(customer_id=cust_normal.id, amount=5.0, is_voided=True)

        # Expenses: 1 Direct, 1 Shared, 1 Farm-Wide
        exp_direct = Expense(amount=8.0, category="Fertilizer", description="=cmd|' /C calc'!A0", scope="specific")
        exp_direct.allocated_units.append(lettuce)

        exp_shared = Expense(amount=10.0, category="Fuel", description="Tractor diesel", scope="shared")
        exp_shared.allocated_units.extend([lettuce, celery])

        exp_farm = Expense(amount=25.0, category="Electricity", description="General farm power", scope="farm")

        db.session.add_all([sale_active, sale_voided, sale_celery, pmt_active, pmt_voided, exp_direct, exp_shared, exp_farm])
        db.session.commit()

        return {
            "lettuce_id": lettuce_id,
            "celery_id": celery_id,
            "cust_normal": cust_normal,
            "cust_malicious": cust_malicious
        }

def test_export_hub_page(client, export_test_data):
    """Verifies the /export hub renders successfully."""
    res = client.get('/export')
    assert res.status_code == 200
    assert b"Export" in res.data
    assert b"Excel" in res.data or b".xlsx" in res.data
    assert b"Sales" in res.data

def test_sales_csv_export(client, export_test_data):
    """Verifies Sales CSV download, column headers, and void status visibility."""
    res = client.get('/export/csv/sales?date_filter=all_time')
    assert res.status_code == 200
    assert res.headers["Content-Type"].startswith("text/csv")
    assert "sales_export_" in res.headers["Content-Disposition"]

    content = res.data.decode('utf-8-sig')
    reader = list(csv.reader(io.StringIO(content)))
    headers = reader[0]

    assert headers == ['Sale ID', 'Date', 'Customer', 'Product / Business Unit', 'Quantity', 'Unit Price ($)', 'Total Amount ($)', 'Status', 'Notes']
    statuses = [row[7] for row in reader[1:]]
    assert "Active" in statuses
    assert "VOIDED" in statuses

def test_expenses_csv_export(client, export_test_data):
    """Verifies Expense CSV handles V1.2 scopes and multi-unit allocations correctly."""
    res = client.get('/export/csv/expenses?date_filter=all_time')
    assert res.status_code == 200
    content = res.data.decode('utf-8-sig')
    reader = list(csv.reader(io.StringIO(content)))

    headers = reader[0]
    assert headers == ['Expense ID', 'Date', 'Category', 'Description', 'Scope', 'Allocated Business Units', 'Amount ($)', 'Status']

    scopes = [row[4] for row in reader[1:]]
    assert "Specific" in scopes
    assert "Shared" in scopes
    assert "Farm" in scopes

def test_customer_balances_all_time_rule(client, export_test_data):
    """Verifies Customer Balances CSV strictly calculates all-time debt regardless of date filters."""
    res = client.get('/export/csv/customer_balances?date_filter=today')
    assert res.status_code == 200
    content = res.data.decode('utf-8-sig')
    reader = list(csv.reader(io.StringIO(content)))

    # Local Grocer bought $20 active ($10 voided ignored), paid $15 active ($5 voided ignored). Balance = $5.00
    grocer_row = next(r for r in reader[1:] if "Local Grocer" in r[1])
    assert grocer_row[5] == "20.00"
    assert grocer_row[6] == "15.00"
    assert grocer_row[7] == "5.00"

def test_business_summary_csv_prevents_double_counting(client, export_test_data):
    """Verifies Business Summary CSV separates direct expenses from shared memo amounts."""
    res = client.get('/export/csv/business_summaries?date_filter=all_time')
    assert res.status_code == 200
    content = res.data.decode('utf-8-sig')
    reader = list(csv.reader(io.StringIO(content)))

    lettuce_row = next(r for r in reader[1:] if "Fancy Lettuce" in r[0])
    assert lettuce_row[1] == "1"
    assert lettuce_row[2] == "20.00"
    assert lettuce_row[3] == "8.00"
    assert lettuce_row[4] == "12.00"
    assert lettuce_row[5] == "60.0"
    assert lettuce_row[6] == "10.00"

def test_formula_injection_sanitization(client, export_test_data):
    """Verifies cells with untrusted leading characters (=, +, -, @) are sanitized with leading apostrophes."""
    res = client.get('/export/csv/sales?date_filter=all_time')
    content = res.data.decode('utf-8-sig')
    reader = list(csv.reader(io.StringIO(content)))
    malicious_row = next(r for r in reader[1:] if "SUM" in r[2])
    assert malicious_row[2].startswith("'=")

def test_excel_workbook_export(client, export_test_data):
    """Verifies Excel .xlsx generation with all 5 required worksheets and correct data types."""
    res = client.get('/export/excel?date_filter=all_time')
    assert res.status_code == 200
    assert res.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "refactored_potato_financials_" in res.headers["Content-Disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    expected_sheets = ["Sales", "Expenses", "Payments", "Customer Balances", "Business Unit Summary"]
    assert wb.sheetnames == expected_sheets

    ws_sales = wb["Sales"]
    assert ws_sales.freeze_panes == "A5"
    assert ws_sales.cell(row=1, column=1).value == "Refactored Potato — Sales"

def test_unit_filtering_on_export(client, export_test_data):
    """Verifies unit_id query parameter isolates records to the selected business unit."""
    lettuce_id = export_test_data["lettuce_id"]
    res = client.get(f'/export/csv/sales?date_filter=all_time&unit_id={lettuce_id}')
    content = res.data.decode('utf-8-sig')
    reader = list(csv.reader(io.StringIO(content)))

    units = [r[3] for r in reader[1:]]
    assert all("Fancy Lettuce" in u for u in units)
    assert not any("Celery" in u for u in units)

def test_empty_dataset_exports(client, app):
    """Verifies exports handle zero records gracefully without 500 errors."""
    with app.app_context():
        res_csv = client.get('/export/csv/sales?date_filter=today')
        assert res_csv.status_code == 200

        res_xlsx = client.get('/export/excel?date_filter=today')
        assert res_xlsx.status_code == 200

def test_export_does_not_modify_database(client, export_test_data):
    """Verifies exporting operations are strictly read-only."""
    sales_before = Sale.query.count()
    expenses_before = Expense.query.count()

    client.get('/export/excel?date_filter=all_time')
    client.get('/export/csv/sales')
    client.get('/export/csv/expenses')

    assert Sale.query.count() == sales_before
    assert Expense.query.count() == expenses_before